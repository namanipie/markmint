import fitz
import docx
from pathlib import Path

base_dir = Path("data/downloads/cycle_tests")

print("=== Testing Text Extraction on CT PDFs and DOCX ===")

samples = [
    ("Data Structures And Algorithm", "CT2_DSA CT 2 SET A Key.pdf"),
    ("Data Structures And Algorithm", "CT1_DSA CT 1 - Ans Key.docx"),
    ("Computer Organization And Architecture", "CT3_coa ct3 set a answer key.pdf"),
    ("Programming For Problem Solving", "CT1_pps ct 1 set 1.pdf"),
    ("Database Management Systems", "CT1_BATCH 1 set1.pdf"),
    ("Design And Analysis Of Algorithms", "CT1_DAA CT1 SET C Answer Key - ques papter ct 1.pdf"),
    ("Chemistry", "CT1_CHEMISTRY CT 1.pdf"),
    ("Calculus And Linear Algebra", "CT1_Updated Answer Key B2 FT2 Set A.pdf")
]

for subj, fname in samples:
    fpath = base_dir / subj / fname
    print(f"\n----------------------------------------")
    print(f"Subject: {subj}")
    print(f"File: {fname} (Exists: {fpath.exists()})")
    if not fpath.exists():
        continue

    text = ""
    if fpath.suffix.lower() == ".pdf":
        try:
            doc = fitz.open(fpath)
            print(f"Pages: {len(doc)}")
            for page in doc:
                text += page.get_text() + "\n"
        except Exception as e:
            print(f"Error reading PDF: {e}")
    elif fpath.suffix.lower() == ".docx":
        try:
            doc = docx.Document(fpath)
            for p in doc.paragraphs:
                text += p.text + "\n"
            for tbl in doc.tables:
                for row in tbl.rows:
                    text += " | ".join(c.text.strip() for c in row.cells) + "\n"
        except Exception as e:
            print(f"Error reading DOCX: {e}")

    print(f"Extracted characters: {len(text)}")
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    print(f"Sample first 10 lines:")
    for l in lines[:10]:
        print(f"  > {l[:100]}")
