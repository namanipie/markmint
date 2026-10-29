import fitz
import docx
import json
import re
from pathlib import Path

def clean_text(text: str) -> str:
    # Normalize spaces and dashes
    t = text.replace('\u2013', '-').replace('\u2014', '-').replace('\u2019', "'").replace('\u201c', '"').replace('\u201d', '"')
    t = re.sub(r'[ \t]+', ' ', t)
    return t.strip()

def parse_docx_questions(fpath: Path, subject: str, cycle: str):
    questions = []
    try:
        doc = docx.Document(fpath)
        full_text = []
        for p in doc.paragraphs:
            if p.text.strip():
                full_text.append(p.text.strip())
        for tbl in doc.tables:
            for row in tbl.rows:
                row_str = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if row_str:
                    full_text.append(row_str)

        text_blob = "\n".join(full_text)
        
        # Regex to match question starts like: "1. ", "Q1. ", "1) ", "Question 1:"
        q_blocks = re.split(r'\n(?=(?:Q(?:uestion)?\s*\.?\s*\d+|\b\d{1,2}\s*[\.\)\-]))', text_blob)
        
        for blk in q_blocks:
            blk = clean_text(blk)
            if len(blk) < 15 or len(blk) > 2500:
                continue
            
            # Check if block looks like a genuine question
            m = re.match(r'^(?:Q(?:uestion)?\s*\.?\s*)?(\d{1,2})\s*[\.\)\-]\s*(.+)', blk, re.DOTALL)
            if not m:
                continue
            
            q_num = int(m.group(1))
            body = m.group(2).strip()

            # Ignore table headers or administrative metadata
            if any(header in body.lower() for header in ['course outcome', 'program specific', 'academic year', 'max. marks', 'duration:']):
                continue

            # Detect marks
            marks = 2
            marks_m = re.search(r'\[?\(?(\d{1,2})\s*(?:Marks?|M)\)?\]?', blk, re.IGNORECASE)
            if marks_m:
                try:
                    marks = int(marks_m.group(1))
                except:
                    pass
            elif 'mcq' in blk.lower() or re.search(r'[a-d]\)\s+[^\n]+[a-d]\)', blk):
                marks = 1

            # Detect question type
            q_type = "DESCRIPTIVE"
            if re.search(r'(?i)\b(?:a\)|[A-D]\.|\(a\)|\(A\))\b', body):
                q_type = "MCQ"
            elif marks <= 2:
                q_type = "SHORT_ANSWER"
            elif any(k in body.lower() for k in ['algorithm', 'pseudocode', 'code', 'program', 'function', 'class']):
                q_type = "CODING"
            elif any(k in body.lower() for k in ['calculate', 'derive', 'evaluate', 'compute', 'find', 'solve']):
                q_type = "NUMERICAL"

            # Check for answer key in block
            ans_m = re.search(r'(?i)(?:Ans(?:wer)?|Key|Solution)\s*[:\-]?\s*([^\n]+)', blk)
            answer_key = ans_m.group(1).strip() if ans_m else None

            questions.append({
                'q_num': q_num,
                'question_text': body,
                'marks': marks,
                'q_type': q_type,
                'answer_key': answer_key,
                'subject': subject,
                'cycle': cycle,
                'source_file': fpath.name
            })
    except Exception as e:
        pass
        
    return questions

def parse_pdf_questions(fpath: Path, subject: str, cycle: str):
    questions = []
    try:
        doc = fitz.open(fpath)
        pages_text = []
        for p in doc:
            t = p.get_text()
            if t.strip():
                pages_text.append(t)
        
        full_text = "\n".join(pages_text)
        
        # Regex to split on question boundaries
        q_blocks = re.split(r'\n(?=(?:Q(?:uestion)?\s*\.?\s*\d+|\b\d{1,2}\s*[\.\)\-]))', full_text)
        
        for blk in q_blocks:
            blk = clean_text(blk)
            if len(blk) < 20 or len(blk) > 3000:
                continue

            m = re.match(r'^(?:Q(?:uestion)?\s*\.?\s*)?(\d{1,2})\s*[\.\)\-]\s*(.+)', blk, re.DOTALL)
            if not m:
                continue

            q_num = int(m.group(1))
            body = m.group(2).strip()

            if any(header in body.lower() for header in ['course outcome', 'program specific', 'academic year', 'max. marks', 'duration:', 'college of engineering']):
                continue

            # Detect marks
            marks = 2
            marks_m = re.search(r'\[?\(?(\d{1,2})\s*(?:Marks?|M)\)?\]?', blk, re.IGNORECASE)
            if marks_m:
                try:
                    marks = int(marks_m.group(1))
                except:
                    pass
            elif re.search(r'(?i)\b(?:a\)|[A-D]\.|\(a\)|\(A\))\b', body):
                marks = 1

            # Detect type
            q_type = "DESCRIPTIVE"
            if re.search(r'(?i)\b(?:a\)|[A-D]\.|\(a\)|\(A\))\b', body):
                q_type = "MCQ"
            elif marks <= 2:
                q_type = "SHORT_ANSWER"
            elif any(k in body.lower() for k in ['algorithm', 'pseudocode', 'code', 'program', 'function', 'class']):
                q_type = "CODING"
            elif any(k in body.lower() for k in ['calculate', 'derive', 'evaluate', 'compute', 'find', 'solve']):
                q_type = "NUMERICAL"

            # Check for answer key in block
            ans_m = re.search(r'(?i)(?:Ans(?:wer)?|Key|Solution)\s*[:\-]?\s*([^\n]+)', blk)
            answer_key = ans_m.group(1).strip() if ans_m else None

            questions.append({
                'q_num': q_num,
                'question_text': body,
                'marks': marks,
                'q_type': q_type,
                'answer_key': answer_key,
                'subject': subject,
                'cycle': cycle,
                'source_file': fpath.name
            })
    except Exception as e:
        pass

    return questions

def main():
    print("==================================================")
    print("MarkMint: Cycle Test Question Parsing Pipeline")
    print("==================================================")

    base_dir = Path("data/downloads/cycle_tests")
    subjs = [d for d in base_dir.iterdir() if d.is_dir()]
    print(f"Scanning {len(subjs)} subject directories...")

    all_extracted_questions = []

    for sdir in sorted(subjs):
        subj_name = sdir.name
        files = list(sdir.glob("*.*"))
        subj_questions = []

        for f in files:
            # Determine cycle
            cycle = "CT"
            fname_low = f.name.lower()
            if 'ct1' in fname_low:
                cycle = "CT1"
            elif 'ct2' in fname_low:
                cycle = "CT2"
            elif 'ct3' in fname_low:
                cycle = "CT3"
            elif 'model' in fname_low:
                cycle = "MODEL"

            if f.suffix.lower() == ".pdf":
                qs = parse_pdf_questions(f, subj_name, cycle)
                subj_questions.extend(qs)
            elif f.suffix.lower() == ".docx":
                qs = parse_docx_questions(f, subj_name, cycle)
                subj_questions.extend(qs)

        # Deduplicate questions within subject by text similarity
        seen_texts = set()
        deduped_subj_qs = []
        for q in subj_questions:
            # First 50 chars of text normalized
            key = re.sub(r'[^a-zA-Z0-9]', '', q['question_text'][:50].lower())
            if key and key not in seen_texts:
                seen_texts.add(key)
                deduped_subj_qs.append(q)

        all_extracted_questions.extend(deduped_subj_qs)
        print(f"  [{subj_name}] Extracted {len(deduped_subj_qs)} unique Cycle Test questions from {len(files)} files")

    print("\n==================================================")
    print(f"Total Extracted Cycle Test Questions: {len(all_extracted_questions)}")
    print("==================================================")

    # Question types breakdown
    type_counts = {}
    cycle_counts = {}
    for q in all_extracted_questions:
        t = q['q_type']
        c = q['cycle']
        type_counts[t] = type_counts.get(t, 0) + 1
        cycle_counts[c] = cycle_counts.get(c, 0) + 1

    print("\nQuestion Breakdown by Type:")
    for t, cnt in sorted(type_counts.items(), key=lambda x: -x[1]):
        print(f"  {t}: {cnt}")

    print("\nQuestion Breakdown by Cycle:")
    for c, cnt in sorted(cycle_counts.items(), key=lambda x: -x[1]):
        print(f"  {c}: {cnt}")

    out_file = Path("data/extracted_cycle_test_questions.json")
    with open(out_file, 'w', encoding='utf-8') as f:
        json.dump(all_extracted_questions, f, indent=2, ensure_ascii=False)

    print(f"\nSaved master question dataset to {out_file}")

if __name__ == '__main__':
    main()
